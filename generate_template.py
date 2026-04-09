from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
import pptx.oxml.ns as nsmap
from lxml import etree

# ── Palette ──────────────────────────────────────────────────────────────────
NAVY       = RGBColor(0x0D, 0x1B, 0x2A)   # slide background / headers
STEEL      = RGBColor(0x1B, 0x3A, 0x5C)   # section bars
SILVER     = RGBColor(0xE8, 0xEC, 0xF0)   # body background
ACCENT     = RGBColor(0x2E, 0x86, 0xC1)   # accent line / bullets
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GREY = RGBColor(0xB0, 0xBE, 0xC5)
DARK_GREY  = RGBColor(0x2C, 0x3E, 0x50)

W = Inches(13.33)   # widescreen 16:9
H = Inches(7.5)

prs = Presentation()
prs.slide_width  = W
prs.slide_height = H

blank = prs.slide_layouts[6]   # completely blank layout


def rgb(r, g, b):
    return RGBColor(r, g, b)


def add_rect(slide, left, top, width, height, fill_color, alpha=None):
    shape = slide.shapes.add_shape(1, left, top, width, height)
    shape.line.fill.background()
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    return shape


def add_text(slide, text, left, top, width, height,
             font_size=18, bold=False, color=WHITE,
             align=PP_ALIGN.LEFT, wrap=True, italic=False):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = "Calibri"
    return txBox


def add_divider(slide, left, top, width, height=Pt(2), color=ACCENT):
    line = slide.shapes.add_shape(1, left, top, width, int(height))
    line.fill.solid()
    line.fill.fore_color.rgb = color
    line.line.fill.background()
    return line


# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 1 — Title / Cover
# ─────────────────────────────────────────────────────────────────────────────
slide = prs.slides.add_slide(blank)

# Full navy background
add_rect(slide, 0, 0, W, H, NAVY)

# Left accent bar
add_rect(slide, 0, 0, Inches(0.18), H, ACCENT)

# Bottom accent strip
add_rect(slide, 0, H - Inches(0.12), W, Inches(0.12), ACCENT)

# Company / Department label
add_text(slide, "ENGINEERING  ·  TECHNOLOGY  ·  INFRASTRUCTURE",
         Inches(0.5), Inches(0.5), Inches(10), Inches(0.5),
         font_size=9, bold=False, color=LIGHT_GREY, italic=True)

# Title
add_text(slide, "Technical Presentation Title",
         Inches(0.5), Inches(2.0), Inches(10), Inches(1.6),
         font_size=40, bold=True, color=WHITE)

# Accent line under title
add_divider(slide, Inches(0.5), Inches(3.7), Inches(6), Pt(3))

# Subtitle
add_text(slide, "Strategic Overview & Recommendations",
         Inches(0.5), Inches(3.9), Inches(9), Inches(0.8),
         font_size=18, bold=False, color=LIGHT_GREY)

# Meta info
add_text(slide, "Presenter Name  |  Role / Team\nDate  |  Confidential",
         Inches(0.5), Inches(6.0), Inches(8), Inches(1.0),
         font_size=12, bold=False, color=LIGHT_GREY)


# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 2 — Agenda
# ─────────────────────────────────────────────────────────────────────────────
slide = prs.slides.add_slide(blank)
add_rect(slide, 0, 0, W, H, SILVER)
add_rect(slide, 0, 0, W, Inches(1.4), NAVY)
add_rect(slide, 0, 0, Inches(0.1), H, ACCENT)

add_text(slide, "AGENDA", Inches(0.35), Inches(0.3), Inches(10), Inches(0.8),
         font_size=28, bold=True, color=WHITE)

items = [
    ("01", "Executive Summary"),
    ("02", "Problem Statement & Scope"),
    ("03", "Technical Deep Dive"),
    ("04", "Architecture / Design"),
    ("05", "Metrics & Results"),
    ("06", "Risks & Mitigations"),
    ("07", "Recommendations & Next Steps"),
]

row_h = Inches(0.72)
start_y = Inches(1.55)

for i, (num, label) in enumerate(items):
    y = start_y + i * row_h
    bg = NAVY if i % 2 == 0 else STEEL
    add_rect(slide, Inches(0.35), y, Inches(12.5), row_h - Inches(0.06), bg)
    add_text(slide, num,   Inches(0.5),  y + Inches(0.12), Inches(0.7),  row_h, font_size=18, bold=True,  color=ACCENT)
    add_text(slide, label, Inches(1.25), y + Inches(0.12), Inches(11.0), row_h, font_size=16, bold=False, color=WHITE)


# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 3 — Executive Summary
# ─────────────────────────────────────────────────────────────────────────────
slide = prs.slides.add_slide(blank)
add_rect(slide, 0, 0, W, H, SILVER)
add_rect(slide, 0, 0, W, Inches(1.4), NAVY)
add_rect(slide, 0, 0, Inches(0.1), H, ACCENT)

add_text(slide, "EXECUTIVE SUMMARY",
         Inches(0.35), Inches(0.3), Inches(10), Inches(0.8),
         font_size=28, bold=True, color=WHITE)
add_text(slide, "Slide 3 of 9",
         Inches(11.5), Inches(0.45), Inches(1.5), Inches(0.5),
         font_size=10, color=LIGHT_GREY, align=PP_ALIGN.RIGHT)

# Three KPI boxes
kpis = [
    ("OBJECTIVE",   "State the business/technical goal clearly and concisely."),
    ("APPROACH",    "Describe the methodology or solution strategy chosen."),
    ("OUTCOME",     "Summarise the expected or achieved result with impact."),
]
box_w = Inches(3.9)
gap   = Inches(0.25)
box_h = Inches(4.5)
start_x = Inches(0.35)

for j, (heading, body) in enumerate(kpis):
    x = start_x + j * (box_w + gap)
    add_rect(slide, x, Inches(1.6), box_w, box_h, WHITE)
    # top colour bar
    add_rect(slide, x, Inches(1.6), box_w, Inches(0.08), ACCENT)
    add_text(slide, heading, x + Inches(0.2), Inches(1.75),
             box_w - Inches(0.4), Inches(0.5),
             font_size=12, bold=True, color=ACCENT)
    add_divider(slide, x + Inches(0.2), Inches(2.3),
                box_w - Inches(0.4), Pt(1))
    add_text(slide, body, x + Inches(0.2), Inches(2.4),
             box_w - Inches(0.4), Inches(3.5),
             font_size=13, color=DARK_GREY, wrap=True)

add_text(slide, "Key Takeaway: Replace with one crisp sentence the leadership must remember.",
         Inches(0.35), Inches(6.3), Inches(12.6), Inches(0.9),
         font_size=12, italic=True, color=STEEL, align=PP_ALIGN.CENTER)


# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 4 — Problem Statement
# ─────────────────────────────────────────────────────────────────────────────
slide = prs.slides.add_slide(blank)
add_rect(slide, 0, 0, W, H, SILVER)
add_rect(slide, 0, 0, W, Inches(1.4), NAVY)
add_rect(slide, 0, 0, Inches(0.1), H, ACCENT)

add_text(slide, "PROBLEM STATEMENT & SCOPE",
         Inches(0.35), Inches(0.3), Inches(10), Inches(0.8),
         font_size=28, bold=True, color=WHITE)
add_text(slide, "Slide 4 of 9",
         Inches(11.5), Inches(0.45), Inches(1.5), Inches(0.5),
         font_size=10, color=LIGHT_GREY, align=PP_ALIGN.RIGHT)

# Left column — problem
add_rect(slide, Inches(0.35), Inches(1.6), Inches(6.0), Inches(5.5), WHITE)
add_rect(slide, Inches(0.35), Inches(1.6), Inches(6.0), Inches(0.08), ACCENT)
add_text(slide, "The Challenge",
         Inches(0.55), Inches(1.75), Inches(5.6), Inches(0.5),
         font_size=14, bold=True, color=NAVY)
add_text(slide,
         "• Describe the root cause of the problem\n"
         "• Quantify the business impact (cost, time, quality)\n"
         "• Explain why this must be solved now\n"
         "• Identify affected stakeholders / systems",
         Inches(0.55), Inches(2.4), Inches(5.6), Inches(4.2),
         font_size=13, color=DARK_GREY, wrap=True)

# Right column — scope
add_rect(slide, Inches(6.7), Inches(1.6), Inches(6.0), Inches(5.5), WHITE)
add_rect(slide, Inches(6.7), Inches(1.6), Inches(6.0), Inches(0.08), ACCENT)
add_text(slide, "Scope",
         Inches(6.9), Inches(1.75), Inches(5.6), Inches(0.5),
         font_size=14, bold=True, color=NAVY)
add_text(slide,
         "In Scope\n"
         "  ✓  System A — module X\n"
         "  ✓  Integration with System B\n"
         "  ✓  Phase 1 deployment region\n\n"
         "Out of Scope\n"
         "  ✗  Legacy system migration\n"
         "  ✗  Third-party SLA changes",
         Inches(6.9), Inches(2.4), Inches(5.6), Inches(4.2),
         font_size=13, color=DARK_GREY, wrap=True)


# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 5 — Technical Deep Dive  (bullet + visual placeholder)
# ─────────────────────────────────────────────────────────────────────────────
slide = prs.slides.add_slide(blank)
add_rect(slide, 0, 0, W, H, SILVER)
add_rect(slide, 0, 0, W, Inches(1.4), NAVY)
add_rect(slide, 0, 0, Inches(0.1), H, ACCENT)

add_text(slide, "TECHNICAL DEEP DIVE",
         Inches(0.35), Inches(0.3), Inches(10), Inches(0.8),
         font_size=28, bold=True, color=WHITE)
add_text(slide, "Slide 5 of 9",
         Inches(11.5), Inches(0.45), Inches(1.5), Inches(0.5),
         font_size=10, color=LIGHT_GREY, align=PP_ALIGN.RIGHT)

add_rect(slide, Inches(0.35), Inches(1.6), Inches(5.8), Inches(5.5), WHITE)
add_rect(slide, Inches(0.35), Inches(1.6), Inches(5.8), Inches(0.08), ACCENT)
add_text(slide, "Key Technical Points",
         Inches(0.55), Inches(1.75), Inches(5.4), Inches(0.5),
         font_size=14, bold=True, color=NAVY)
add_text(slide,
         "• Component / layer description\n\n"
         "• Technology choices & rationale\n\n"
         "• Data flow / processing logic\n\n"
         "• Performance characteristics\n\n"
         "• Security & compliance considerations",
         Inches(0.55), Inches(2.4), Inches(5.4), Inches(4.2),
         font_size=13, color=DARK_GREY, wrap=True)

# Diagram placeholder
add_rect(slide, Inches(6.4), Inches(1.6), Inches(6.55), Inches(5.5), WHITE)
add_rect(slide, Inches(6.4), Inches(1.6), Inches(6.55), Inches(0.08), ACCENT)
add_text(slide, "[ Insert Architecture / Flow Diagram ]",
         Inches(6.6), Inches(3.8), Inches(6.1), Inches(1.0),
         font_size=14, color=LIGHT_GREY, align=PP_ALIGN.CENTER)


# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 6 — Metrics & Results
# ─────────────────────────────────────────────────────────────────────────────
slide = prs.slides.add_slide(blank)
add_rect(slide, 0, 0, W, H, SILVER)
add_rect(slide, 0, 0, W, Inches(1.4), NAVY)
add_rect(slide, 0, 0, Inches(0.1), H, ACCENT)

add_text(slide, "METRICS & RESULTS",
         Inches(0.35), Inches(0.3), Inches(10), Inches(0.8),
         font_size=28, bold=True, color=WHITE)
add_text(slide, "Slide 6 of 9",
         Inches(11.5), Inches(0.45), Inches(1.5), Inches(0.5),
         font_size=10, color=LIGHT_GREY, align=PP_ALIGN.RIGHT)

metrics = [
    ("99.95 %",  "Availability SLA"),
    ("< 200 ms", "p99 Latency"),
    ("40 %",     "Cost Reduction"),
    ("3×",       "Throughput Gain"),
]
mw = Inches(2.9)
mh = Inches(2.2)
gap = Inches(0.25)
sy  = Inches(1.7)
sx  = Inches(0.35)

for k, (val, lbl) in enumerate(metrics):
    x = sx + k * (mw + gap)
    add_rect(slide, x, sy, mw, mh, NAVY)
    add_text(slide, val, x, sy + Inches(0.35), mw, Inches(0.9),
             font_size=30, bold=True, color=ACCENT, align=PP_ALIGN.CENTER)
    add_text(slide, lbl, x, sy + Inches(1.25), mw, Inches(0.6),
             font_size=12, color=LIGHT_GREY, align=PP_ALIGN.CENTER)

# Chart placeholder
add_rect(slide, Inches(0.35), Inches(4.1), Inches(12.55), Inches(3.0), WHITE)
add_rect(slide, Inches(0.35), Inches(4.1), Inches(12.55), Inches(0.07), ACCENT)
add_text(slide, "[ Insert Chart / Graph — Trend or Comparison ]",
         Inches(0.35), Inches(5.2), Inches(12.55), Inches(0.8),
         font_size=14, color=LIGHT_GREY, align=PP_ALIGN.CENTER)


# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 7 — Risks & Mitigations
# ─────────────────────────────────────────────────────────────────────────────
slide = prs.slides.add_slide(blank)
add_rect(slide, 0, 0, W, H, SILVER)
add_rect(slide, 0, 0, W, Inches(1.4), NAVY)
add_rect(slide, 0, 0, Inches(0.1), H, ACCENT)

add_text(slide, "RISKS & MITIGATIONS",
         Inches(0.35), Inches(0.3), Inches(10), Inches(0.8),
         font_size=28, bold=True, color=WHITE)
add_text(slide, "Slide 7 of 9",
         Inches(11.5), Inches(0.45), Inches(1.5), Inches(0.5),
         font_size=10, color=LIGHT_GREY, align=PP_ALIGN.RIGHT)

# Table header
add_rect(slide, Inches(0.35), Inches(1.55), Inches(12.55), Inches(0.45), NAVY)
for hx, hw, ht in [
    (Inches(0.35), Inches(3.5),  "Risk"),
    (Inches(3.95), Inches(1.5),  "Likelihood"),
    (Inches(5.55), Inches(1.5),  "Impact"),
    (Inches(7.15), Inches(5.75), "Mitigation"),
]:
    add_text(slide, ht, hx + Inches(0.1), Inches(1.6), hw, Inches(0.4),
             font_size=11, bold=True, color=WHITE)

rows = [
    ("Vendor API deprecation",    "Medium", "High",   "Maintain abstraction layer; evaluate alternatives Q3"),
    ("Data pipeline latency spike","Low",   "High",   "Circuit breaker + retry logic already implemented"),
    ("Team capacity during ramp",  "High",  "Medium", "Backfill via contractor; cross-train 2 engineers"),
    ("Regulatory change",          "Low",   "High",   "Monthly compliance review; legal sign-off on design"),
]

SEVERITY_COLOR = {"High": rgb(0xC0, 0x39, 0x2B), "Medium": rgb(0xD3, 0x54, 0x00), "Low": rgb(0x27, 0x7A, 0x27)}

for ri, (risk, lk, imp, mit) in enumerate(rows):
    ry = Inches(2.0) + ri * Inches(1.1)
    bg = WHITE if ri % 2 == 0 else rgb(0xF2, 0xF5, 0xF7)
    add_rect(slide, Inches(0.35), ry, Inches(12.55), Inches(1.05), bg)
    add_text(slide, risk, Inches(0.5),  ry + Inches(0.12), Inches(3.3),  Inches(0.9), font_size=12, color=DARK_GREY)
    # Likelihood badge
    lk_box = slide.shapes.add_shape(1, Inches(3.95) + Inches(0.1), ry + Inches(0.2), Inches(1.1), Inches(0.55))
    lk_box.fill.solid(); lk_box.fill.fore_color.rgb = SEVERITY_COLOR.get(lk, DARK_GREY)
    lk_box.line.fill.background()
    add_text(slide, lk,  Inches(3.95) + Inches(0.1), ry + Inches(0.22), Inches(1.1), Inches(0.55), font_size=10, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    imp_box = slide.shapes.add_shape(1, Inches(5.55) + Inches(0.1), ry + Inches(0.2), Inches(1.1), Inches(0.55))
    imp_box.fill.solid(); imp_box.fill.fore_color.rgb = SEVERITY_COLOR.get(imp, DARK_GREY)
    imp_box.line.fill.background()
    add_text(slide, imp, Inches(5.55) + Inches(0.1), ry + Inches(0.22), Inches(1.1), Inches(0.55), font_size=10, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_text(slide, mit, Inches(7.25), ry + Inches(0.12), Inches(5.5), Inches(0.9), font_size=11, color=DARK_GREY, wrap=True)


# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 8 — Recommendations & Next Steps
# ─────────────────────────────────────────────────────────────────────────────
slide = prs.slides.add_slide(blank)
add_rect(slide, 0, 0, W, H, SILVER)
add_rect(slide, 0, 0, W, Inches(1.4), NAVY)
add_rect(slide, 0, 0, Inches(0.1), H, ACCENT)

add_text(slide, "RECOMMENDATIONS & NEXT STEPS",
         Inches(0.35), Inches(0.3), Inches(11), Inches(0.8),
         font_size=26, bold=True, color=WHITE)
add_text(slide, "Slide 8 of 9",
         Inches(11.5), Inches(0.45), Inches(1.5), Inches(0.5),
         font_size=10, color=LIGHT_GREY, align=PP_ALIGN.RIGHT)

# Recommendations left
add_rect(slide, Inches(0.35), Inches(1.6), Inches(5.8), Inches(5.5), WHITE)
add_rect(slide, Inches(0.35), Inches(1.6), Inches(5.8), Inches(0.08), ACCENT)
add_text(slide, "Recommendations",
         Inches(0.55), Inches(1.75), Inches(5.4), Inches(0.5),
         font_size=14, bold=True, color=NAVY)
recs = [
    ("R1", "Approve Phase 1 rollout to production — low risk, high ROI"),
    ("R2", "Allocate budget for observability tooling upgrade"),
    ("R3", "Initiate vendor renegotiation for Q3 contract renewal"),
]
for ri, (tag, txt) in enumerate(recs):
    ry = Inches(2.4) + ri * Inches(1.4)
    add_rect(slide, Inches(0.55), ry, Inches(0.55), Inches(0.55), ACCENT)
    add_text(slide, tag, Inches(0.55), ry + Inches(0.05), Inches(0.55), Inches(0.45),
             font_size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_text(slide, txt, Inches(1.2), ry, Inches(4.7), Inches(0.7),
             font_size=12, color=DARK_GREY, wrap=True)

# Timeline right
add_rect(slide, Inches(6.4), Inches(1.6), Inches(6.55), Inches(5.5), WHITE)
add_rect(slide, Inches(6.4), Inches(1.6), Inches(6.55), Inches(0.08), ACCENT)
add_text(slide, "Delivery Timeline",
         Inches(6.6), Inches(1.75), Inches(6.1), Inches(0.5),
         font_size=14, bold=True, color=NAVY)

phases = [
    ("Week 1–2",  "Finalise architecture sign-off & environment setup"),
    ("Week 3–6",  "Development & unit testing"),
    ("Week 7–8",  "Integration testing & staging deployment"),
    ("Week 9",    "Production rollout — phased (10 % → 50 % → 100 %)"),
    ("Week 10",   "Hypercare & KPI review"),
]
for pi, (period, task) in enumerate(phases):
    py = Inches(2.3) + pi * Inches(0.96)
    add_rect(slide, Inches(6.55), py, Inches(1.6), Inches(0.7), NAVY)
    add_text(slide, period, Inches(6.55), py + Inches(0.1), Inches(1.6), Inches(0.55),
             font_size=9, bold=True, color=ACCENT, align=PP_ALIGN.CENTER)
    add_divider(slide, Inches(8.25), py + Inches(0.35), Inches(0.2), Pt(2))
    add_text(slide, task, Inches(8.55), py, Inches(4.2), Inches(0.8),
             font_size=11, color=DARK_GREY, wrap=True)


# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 9 — Closing / Thank You
# ─────────────────────────────────────────────────────────────────────────────
slide = prs.slides.add_slide(blank)
add_rect(slide, 0, 0, W, H, NAVY)
add_rect(slide, 0, 0, Inches(0.18), H, ACCENT)
add_rect(slide, 0, H - Inches(0.12), W, Inches(0.12), ACCENT)

add_text(slide, "Thank You",
         Inches(0.5), Inches(2.2), Inches(12), Inches(1.4),
         font_size=48, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

add_divider(slide, Inches(4.0), Inches(3.7), Inches(5.3), Pt(2))

add_text(slide, "Questions & Discussion",
         Inches(0.5), Inches(3.9), Inches(12), Inches(0.7),
         font_size=20, color=LIGHT_GREY, align=PP_ALIGN.CENTER)

add_text(slide,
         "presenter@company.com   ·   +1 (555) 000-0000   ·   Team / Channel",
         Inches(0.5), Inches(5.4), Inches(12), Inches(0.6),
         font_size=12, color=LIGHT_GREY, align=PP_ALIGN.CENTER, italic=True)

add_text(slide, "CONFIDENTIAL — FOR INTERNAL USE ONLY",
         Inches(0.5), Inches(6.8), Inches(12), Inches(0.5),
         font_size=9, color=LIGHT_GREY, align=PP_ALIGN.CENTER)


# ─────────────────────────────────────────────────────────────────────────────
out = "/home/user/test1/management_presentation_template.pptx"
prs.save(out)
print(f"Saved: {out}")
